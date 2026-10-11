"""Excel reader with strict measured-data and explicit legacy compatibility modes."""

from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook

from .evaluation import DischargeSeries
from .metrics import RATE_ORDER
from .runstore import sha256_file

RATE_C = {"0.5C": 0.5, "1C": 1.0, "2C": 2.0, "3C": 3.0}


@dataclass(frozen=True)
class WorkbookSchema:
    sheet_by_rate: Mapping[str, str]
    time_column: str | None
    current_column: str | None
    voltage_column: str
    temperature_column: str | None = None
    legacy_assumed_1s: bool = False
    sample_period_s: float | None = None
    nominal_capacity_ah: float | None = None
    discharge_current_sign: int | None = None

    def __post_init__(self) -> None:
        if set(self.sheet_by_rate) != set(RATE_ORDER):
            raise ValueError(f"sheet_by_rate must specify {RATE_ORDER}")
        if self.legacy_assumed_1s:
            if self.discharge_current_sign not in {-1, 1}:
                raise ValueError("legacy current sign must be explicitly -1 or +1")
            if self.sample_period_s != 1.0:
                raise ValueError("legacy_assumed_1s requires an explicit sample_period_s=1.0")
            if self.nominal_capacity_ah is None or not math.isfinite(self.nominal_capacity_ah) or self.nominal_capacity_ah <= 0:
                raise ValueError("legacy mode requires a finite positive nominal_capacity_ah")
            if self.time_column is None and (self.sample_period_s is None or self.sample_period_s <= 0):
                raise ValueError("legacy time requires a positive sample_period_s")
            if self.current_column is None and (self.nominal_capacity_ah is None or self.nominal_capacity_ah <= 0):
                raise ValueError("legacy current requires an explicit nominal_capacity_ah")
            if self.current_column is None and self.discharge_current_sign not in {-1, 1}:
                raise ValueError("legacy current sign must be explicitly -1 or +1")


def _column_indexes(headers) -> dict[str, int]:
    result = {}
    for index, header in enumerate(headers):
        if header is None:
            continue
        normalized = str(header).strip().casefold()
        if normalized in result:
            raise ValueError(f"duplicate column header: {header}")
        result[normalized] = index
    return result


def _read_numeric_column(rows, indexes: dict[str, int], name: str, sheet: str) -> tuple[float, ...]:
    key = name.strip().casefold()
    if key not in indexes:
        raise ValueError(f"sheet {sheet!r} is missing required column {name!r}")
    values = []
    for row_number, row in enumerate(rows, start=2):
        value = row[indexes[key]] if indexes[key] < len(row) else None
        if value is None or isinstance(value, bool):
            raise ValueError(f"sheet {sheet!r}, row {row_number}: {name!r} is blank")
        try:
            number = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"sheet {sheet!r}, row {row_number}: {name!r} is not numeric") from exc
        values.append(number)
    return tuple(values)


def load_four_rate_workbook(path: str | Path, schema: WorkbookSchema) -> tuple[dict[str, DischargeSeries], str]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    digest = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing = set(schema.sheet_by_rate.values()) - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"workbook is missing configured sheets: {sorted(missing)}")
        output = {}
        for rate in RATE_ORDER:
            sheet_name = schema.sheet_by_rate[rate]
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers:
                raise ValueError(f"sheet {sheet_name!r} has no header row")
            rows = list(iterator)
            indexes = _column_indexes(headers)
            voltage = _read_numeric_column(rows, indexes, schema.voltage_column, sheet_name)

            if schema.time_column is not None and schema.time_column.strip().casefold() in indexes:
                time_s = _read_numeric_column(rows, indexes, schema.time_column, sheet_name)
                time_source = "measured"
            elif schema.legacy_assumed_1s and schema.sample_period_s is not None:
                time_s = tuple(i * schema.sample_period_s for i in range(len(rows)))
                time_source = "legacy_assumed_1s"
            else:
                raise ValueError(
                    f"sheet {sheet_name!r} has no configured time column; "
                    "enable legacy_assumed_1s explicitly to use row numbers"
                )

            current = None
            current_source = "measured"
            if schema.current_column is not None and schema.current_column.strip().casefold() in indexes:
                column_index = indexes[schema.current_column.strip().casefold()]
                entirely_blank = all(row[column_index] is None for row in rows)
                if not entirely_blank or not schema.legacy_assumed_1s:
                    current = _read_numeric_column(rows, indexes, schema.current_column, sheet_name)
            if current is None:
                if not schema.legacy_assumed_1s or schema.nominal_capacity_ah is None:
                    raise ValueError(
                        f"sheet {sheet_name!r} has no usable measured current; "
                        "legacy_assumed_1s requires an explicit nominal capacity and current sign"
                    )
                current = tuple(
                    schema.discharge_current_sign * RATE_C[rate] * schema.nominal_capacity_ah
                    for _ in rows
                )
                current_source = "legacy_assumed_1s"

            temperature = None
            if schema.temperature_column is not None:
                temperature = _read_numeric_column(rows, indexes, schema.temperature_column, sheet_name)
            assumptions = []
            if time_source == "legacy_assumed_1s":
                assumptions.append(f"sample_period_s={schema.sample_period_s:g}")
            if current_source == "legacy_assumed_1s":
                assumptions.extend((
                    f"nominal_capacity_ah={schema.nominal_capacity_ah:g}",
                    f"discharge_current_sign={schema.discharge_current_sign:+d}",
                    "constant current derived from nominal C-rate",
                ))
            output[rate] = DischargeSeries(
                rate, time_s, voltage, current, temperature, current_source, tuple(assumptions),
                time_source, schema.discharge_current_sign,
            )
        return output, digest
    finally:
        workbook.close()


def inspect_workbook(path: str | Path) -> dict:
    """Read-only structure and numeric completeness report; makes no schema assumptions."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    digest = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None) or ()
            stats = [{"header": header, "nonblank_count": 0, "numeric_count": 0,
                      "numeric_min": None, "numeric_max": None} for header in headers]
            row_count = 1 if headers else 0
            for row in iterator:
                row_count += 1
                for column_index, stat in enumerate(stats):
                    value = row[column_index] if column_index < len(row) else None
                    if value is None or isinstance(value, str) and not value.strip():
                        continue
                    stat["nonblank_count"] += 1
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        number = float(value)
                        stat["numeric_count"] += 1
                        stat["numeric_min"] = number if stat["numeric_min"] is None else min(stat["numeric_min"], number)
                        stat["numeric_max"] = number if stat["numeric_max"] is None else max(stat["numeric_max"], number)
            sheets.append({"name": sheet.title, "row_count_including_header": row_count, "columns": stats})
        return {"path": str(path.resolve()), "sha256": digest, "sheets": sheets}
    finally:
        workbook.close()
